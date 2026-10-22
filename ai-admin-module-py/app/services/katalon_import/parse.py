"""Readers for each kind of Katalon export. Pure functions: text/bytes in, dicts out.

Secrets never leave this module: the execution profile (.glbl) is reduced to
variable names, kinds and non-secret values (URLs, hosts, settings) as it is
read. Passwords, card numbers and user names are not returned at all.
"""
from __future__ import annotations

import csv
import io
import re
import urllib.parse
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from typing import Any, Optional

# --- Converter Markdown (the team's "flattened test case" export) -------------------

_NUMBERED = re.compile(r"^(\d+(?:\.\d+)*)\.\s+(.*)$")
_CALL = re.compile(r"^(.*?)\s{2,}\((test case|keyword)\s+([^)]+)\)\s*$")
_NOT_RUN = re.compile(r"^\(Not run with this test data:\s*(.*?)\s+-\s+(\d+)\s+steps?\s+skipped\)\s*$")
_BACKEND = re.compile(r"^(?:->|- Backend:)\s*(.*)$")
_INPUTS = re.compile(r"^inputs:\s*(.*)$")


@dataclass
class Node:
    """One line of the Steps block, with the lines nested under it."""
    kind: str                      # step | control | not_run
    text: str
    indent: int
    line: int
    number: Optional[str] = None   # "16.63.4" for numbered steps
    call: Optional[dict] = None    # {"type": "test case"|"keyword", "ref": "..."} when the step calls one
    inputs: list[dict] = field(default_factory=list)
    backend: Optional[dict] = None
    children: list["Node"] = field(default_factory=list)

    def as_dict(self) -> dict:
        return {
            "kind": self.kind, "text": self.text, "number": self.number, "line": self.line,
            "call": self.call, "inputs": self.inputs, "backend": self.backend,
            "children": [c.as_dict() for c in self.children],
        }


def _split_inputs(text: str) -> list[dict]:
    """'url = {GlobalVariable.appUrl}; navTab = "USPSalesOrderTab"' -> [{name, raw}]."""
    out = []
    depth, start, parts = 0, 0, []
    for i, ch in enumerate(text):
        if ch in "({[":
            depth += 1
        elif ch in ")}]":
            depth = max(0, depth - 1)
        elif ch == ";" and depth == 0:
            parts.append(text[start:i])
            start = i + 1
    parts.append(text[start:])
    for part in parts:
        if "=" not in part:
            continue
        name, raw = part.split("=", 1)
        out.append({"name": name.strip(), "raw": raw.strip()})
    return out


def _parse_backend(text: str) -> dict:
    """'backend DB/API helper, not expanded  calls: A, B  SQL: s1 | s2' -> parts."""
    calls: list[str] = []
    sql: list[str] = []
    body = text
    m = re.search(r"\bSQL:\s*(.*)$", body)
    if m:
        sql = [s.strip() for s in m.group(1).split(" | ") if s.strip()]
        body = body[: m.start()]
    m = re.search(r"\bcalls:\s*(.*)$", body)
    if m:
        calls = [c.strip() for c in m.group(1).split(",") if c.strip()]
        body = body[: m.start()]
    kind = "prerequisite" if "prerequisite" in body else "helper" if "helper" in body else "backend"
    if "browser/driver" in body:
        kind = "driver"
    return {"kind": kind, "calls": calls, "sql": sql, "note": body.strip(" -")}


def _tree(lines: list[tuple[int, str, int]]) -> list[Node]:
    root = Node(kind="root", text="", indent=-1, line=0)
    stack = [root]
    for indent, text, lineno in lines:
        while stack and indent <= stack[-1].indent:
            stack.pop()
        parent = stack[-1]
        m_in = _INPUTS.match(text)
        if m_in and parent.kind == "step":
            parent.inputs = _split_inputs(m_in.group(1))
            continue
        m_be = _BACKEND.match(text)
        if m_be and parent.kind == "step":
            parent.backend = _parse_backend(m_be.group(1))
            continue
        m_nr = _NOT_RUN.match(text)
        if m_nr:
            parent.children.append(Node(kind="not_run", text=m_nr.group(1).strip(), indent=indent, line=lineno,
                                        number=m_nr.group(2)))
            continue
        m_num = _NUMBERED.match(text)
        if m_num:
            node = Node(kind="step", text=m_num.group(2).strip(), indent=indent, line=lineno, number=m_num.group(1))
            m_call = _CALL.match(node.text)
            if m_call:
                node.text = m_call.group(1).strip()
                node.call = {"type": m_call.group(2), "ref": m_call.group(3).strip()}
        else:
            node = Node(kind="control", text=text.rstrip(":").strip(), indent=indent, line=lineno)
        parent.children.append(node)
        stack.append(node)
    return root.children


def parse_converter_md(text: str) -> dict:
    """The converter's Markdown -> header, data variables, step tree, locators, issues."""
    text = text.replace("\r\n", "\n")
    out: dict[str, Any] = {
        "title": "", "katalon_test_case": None, "script": None, "data_sheet": None, "data_key": None,
        "data_vars": [], "listeners": [], "steps": [], "locators": [], "issues": [],
        "inline_locators": [], "unconverted": [], "format": "v1",
    }
    m = re.search(r"^#\s+(.+)$", text, re.M)
    if m:
        out["title"] = m.group(1).strip()
    m = re.search(r"^- Katalon test case:\s*`([^`]+)`", text, re.M)
    if m:
        out["katalon_test_case"] = m.group(1)
    m = re.search(r"^- Script:\s*`([^`]+)`", text, re.M)
    if m:
        out["script"] = m.group(1)
    m = re.search(r'^- Test data:\s*sheet\s*`([^`]+)`(?:.*?"([^"]+)")?', text, re.M)
    if m:
        out["data_sheet"], out["data_key"] = m.group(1), m.group(2)

    sections = re.split(r"^##\s+", text, flags=re.M)
    for section in sections[1:]:
        heading, _, body = section.partition("\n")
        heading = heading.strip().lower()
        if heading.startswith("test data used"):
            for row in re.findall(r"^\|(.+)\|\s*$", body, re.M):
                cells = [c.strip() for c in row.split("|")]
                if len(cells) >= 3 and cells[0] not in ("Variable", "---") and not set(cells[0]) <= {"-"}:
                    out["data_vars"].append({"name": cells[0], "column": cells[1], "example": cells[2]})
        elif heading.startswith("test listeners"):
            out["listeners"] = [ln.strip() for ln in body.splitlines() if ln.strip()]
        elif heading == "steps":
            block = re.search(r"```text\n(.*?)```", body, re.S)
            lines = []
            for offset, raw in enumerate((block.group(1) if block else "").splitlines()):
                if not raw.strip():
                    continue
                lines.append((len(raw) - len(raw.lstrip(" ")), raw.strip(), offset + 1))
            out["steps"] = [n.as_dict() for n in _tree(lines)]
        elif heading.startswith("source script issues"):
            out["format"] = "v2"
            for ln in body.splitlines():
                m2 = re.match(r"^- (.*?)\s+\(([^()]+\.groovy):(\d+)\)\s*$", ln.strip())
                if m2:
                    out["issues"].append({"text": m2.group(1), "file": m2.group(2), "line": int(m2.group(3))})
        elif heading.startswith("element locators"):
            out["format"] = "v2"
            out["locators"] = _parse_locator_section(body)
        elif heading.startswith("custom selenium"):
            for ln in body.splitlines():
                m2 = re.match(r"^- `(.*)`\s+\(([^()]+):(\d+)\)\s*$", ln.strip())
                if m2:
                    out["inline_locators"].append({
                        "code": m2.group(1), "file": m2.group(2), "line": int(m2.group(3)),
                        "xpaths": re.findall(r'By\.xpath\("((?:[^"\\]|\\.)*)"\)', m2.group(1)),
                    })
        elif "not converted" in heading:
            for ln in body.splitlines():
                m2 = re.match(r"^- `(.*)`\s+\(([^()]+):(\d+)\)\s*$", ln.strip())
                if m2:
                    out["unconverted"].append({"code": m2.group(1), "file": m2.group(2), "line": int(m2.group(3))})
    return out


def _parse_locator_section(body: str) -> list[dict]:
    items = []
    for chunk in re.split(r"^###\s+", body, flags=re.M)[1:]:
        head, _, rest = chunk.partition("\n")
        m = re.match(r"^(.*?)\s+\(`([^`]+)`\)\s*$", head.strip())
        if not m:
            continue
        item: dict[str, Any] = {"label": m.group(1).strip(), "path": m.group(2).strip(), "selectors": [],
                                "active": None, "params": [], "examples": [], "converter_note": None}
        for ln in rest.splitlines():
            s = ln.strip()
            sel = re.match(r"^- (XPATH|BASIC|CSS|IMAGE)(\s*\(active\))?:\s*`(.*)`\s*$", s)
            if sel:
                entry = {"method": sel.group(1), "value": sel.group(3), "active": bool(sel.group(2))}
                item["selectors"].append(entry)
                if entry["active"]:
                    item["active"] = entry
                continue
            dyn = re.match(r"^- DYNAMIC locator:.*Parameters:\s*(.*)$", s)
            if dyn:
                item["params"] = [p.strip(" `") for p in dyn.group(1).split(",") if p.strip(" `")]
                continue
            ex = re.match(r"^- step ([\d.]+):\s*(.*)$", s)
            if ex:
                item["examples"].append({"step": ex.group(1), "text": ex.group(2)})
                continue
            if s.startswith("- Playwright") and "no usable locator" in s:
                item["converter_note"] = "The converter found no Playwright form for this locator."
        if item["active"] is None and item["selectors"]:
            item["active"] = item["selectors"][0]
        items.append(item)
    return items


# --- Object Repository .rs files --------------------------------------------------------

def parse_object_repository(path_in_repo: str, xml_text: str) -> Optional[dict]:
    """One `Object Repository/<path>.rs` file -> the same shape as a locator-section entry."""
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError:
        return None
    if root.tag != "WebElementEntity":
        return None
    method = (root.findtext("selectorMethod") or "BASIC").strip().upper()
    selectors = []
    for entry in root.iter("entry"):
        key, value = entry.findtext("key"), entry.findtext("value")
        if key and value:
            selectors.append({"method": key.strip().upper(), "value": value, "active": key.strip().upper() == method})
    if not any(s["active"] for s in selectors):
        # BASIC objects keep their locator in webElementProperties (xpath / css properties).
        for prop in root.iter("webElementProperties"):
            if (prop.findtext("isSelected") or "").strip() == "true" and prop.findtext("name") in ("xpath", "css"):
                selectors.append({"method": "BASIC", "value": prop.findtext("value") or "", "active": True})
                break
    active = next((s for s in selectors if s["active"]), selectors[0] if selectors else None)
    params = sorted(set(re.findall(r"\$\{(\w+)\}", (active or {}).get("value") or "")))
    path = re.sub(r"^.*?Object Repository/", "", path_in_repo.replace("\\", "/"))
    path = re.sub(r"\.rs$", "", path)
    return {"label": root.findtext("name") or path.rsplit("/", 1)[-1], "path": path, "selectors": selectors,
            "active": active, "params": params, "examples": [], "converter_note": None, "source": "object_repository"}


# --- Test data ----------------------------------------------------------------------------

def parse_katalon_dat(xml_text: str) -> dict:
    """Katalon internal data file -> {name, columns, rows: [{col: value}]}.

    Cells are '+'-encoded and separated by single spaces; the first <data> row
    holds column indexes, not values.
    """
    root = ET.fromstring(xml_text)
    columns = [c.findtext("name") or f"col{i}" for i, c in enumerate(root.findall("internalDataColumns"))]
    rows = []
    for index, d in enumerate(root.findall("data")):
        cells = [urllib.parse.unquote_plus(c) for c in (d.text or "").split(" ")]
        if index == 0 and all(c.isdigit() or c in ("", "null") for c in cells):
            continue
        values = {col: (cells[i] if i < len(cells) and cells[i] != "null" else "") for i, col in enumerate(columns)}
        if any(v for v in values.values()):
            rows.append(values)
    return {"name": root.findtext("name") or "", "columns": columns, "rows": rows}


def parse_csv_data(text: str, name: str) -> dict:
    reader = csv.reader(io.StringIO(text))
    table = [r for r in reader if any(c.strip() for c in r)]
    if not table:
        return {"name": name, "columns": [], "rows": []}
    columns = [c.strip() or f"col{i}" for i, c in enumerate(table[0])]
    rows = [{col: (r[i].strip() if i < len(r) else "") for i, col in enumerate(columns)} for r in table[1:]]
    return {"name": name, "columns": columns, "rows": rows}


def parse_xlsx_data(data: bytes, name: str) -> list[dict]:
    """Every sheet with a header row -> one table each."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c for c in r)]
        if len(rows) < 2:
            continue
        columns = [c or f"col{i}" for i, c in enumerate(rows[0])]
        tables.append({"name": f"{name}:{ws.title}" if len(wb.worksheets) > 1 else name, "columns": columns,
                       "rows": [{col: (r[i] if i < len(r) else "") for i, col in enumerate(columns)} for r in rows[1:]]})
    return tables


# --- Execution profile (.glbl) -- secrets are dropped here ----------------------------

_SECRET_NAME = re.compile(r"pass|pwd|secret|token|credit|card|apikey|api_key|cvv|pin\b", re.I)
_USER_NAME = re.compile(r"user", re.I)
_CONN_NAME = re.compile(r"connection|jdbc|^sqlserver", re.I)
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")


def _strip_quotes(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in "'\"":
        return value[1:-1]
    return value


def _host_of_connection(value: str) -> Optional[dict]:
    v = value
    m = re.search(r"HOST\s*=\s*([\w.\-]+).*?PORT\s*=\s*(\d+).*?(?:SERVICE_NAME|SID)\s*=\s*([\w.\-]+)", v, re.I | re.S)
    if m:
        return {"host": m.group(1), "port": m.group(2), "service": m.group(3), "type": "oracle"}
    m = re.search(r"@(?://)?([\w.\-]+):(\d+)[:/]([\w.\-]+)", v)
    if m:
        return {"host": m.group(1), "port": m.group(2), "service": m.group(3), "type": "oracle"}
    m = re.search(r"jdbc:sqlserver://([\w.\-]+)(?::(\d+))?(?:;.*?databaseName=([\w.\-]+))?", v, re.I)
    if m:
        return {"host": m.group(1), "port": m.group(2) or "1433", "service": m.group(3), "type": "sqlserver"}
    return None


def parse_profile(xml_text: str, name: str) -> dict:
    """{name, variables: [{name, kind, value?, connection?}]}. Never returns a secret or user name."""
    root = ET.fromstring(xml_text)
    variables = []
    for v in root.iter("GlobalVariableEntity"):
        var = (v.findtext("name") or "").strip()
        raw = _strip_quotes(v.findtext("initValue") or "")
        entry: dict[str, Any] = {"name": var, "description": (v.findtext("description") or "").strip()}
        if _SECRET_NAME.search(var):
            entry.update(kind="secret", set=bool(raw))
        elif _USER_NAME.search(var) and not re.search(r"url|path", var, re.I):
            entry.update(kind="user", set=bool(raw))
        elif _CONN_NAME.search(var) or raw.lower().startswith("jdbc:"):
            entry.update(kind="connection", connection=_host_of_connection(raw), set=bool(raw))
        elif re.match(r"^https?://", raw):
            entry.update(kind="url", value=raw)
        elif _EMAIL.search(raw):
            entry.update(kind="email", value=_EMAIL.sub("<email>", raw))
        else:
            entry.update(kind="setting", value=raw[:200])
        variables.append(entry)
    return {"name": name, "variables": variables}


# --- Keyword sources (.groovy) ----------------------------------------------------------

_METHOD = re.compile(
    r"(?P<kw>@Keyword\s+)?(?:public\s+|private\s+)?(?:def|static|String|boolean|void|int|Map\S*|List\S*)"
    r"(?:\s+(?:static|def|String|boolean|void|int|File))*\s+(?P<name>[A-Za-z_]\w*)\s*\((?P<args>[^)]*)\)\s*\{",
)
_SQL_START = re.compile(r"^\s*(select|update|insert|delete|merge)\b", re.I)


def _strings(body: str) -> list[str]:
    found = re.findall(r'"""(.*?)"""', body, re.S)
    rest = re.sub(r'"""(.*?)"""', "", body, flags=re.S)
    found += re.findall(r'"((?:[^"\\\n]|\\.)*)"', rest)
    found += re.findall(r"'((?:[^'\\\n]|\\.){12,})'", rest)
    return found


def parse_groovy_keywords(rel_path: str, text: str) -> dict:
    """One keyword class -> its methods with the SQL they run and the database they open."""
    package = re.search(r"^\s*package\s+([\w.]+)", text, re.M)
    cls = re.search(r"^\s*(?:public\s+)?class\s+(\w+)", text, re.M)
    class_name = cls.group(1) if cls else re.sub(r"\.groovy$", "", rel_path.replace("\\", "/").rsplit("/", 1)[-1])
    qualified = f"{package.group(1)}.{class_name}" if package else class_name
    starts = [m for m in _METHOD.finditer(text) if m.group("name") not in ("if", "for", "while", "catch", "switch")]
    names = {m.group("name") for m in starts}
    methods = []
    for i, m in enumerate(starts):
        body = text[m.end(): starts[i + 1].start() if i + 1 < len(starts) else len(text)]
        statements = []
        for s in _strings(body):
            flat = re.sub(r"\s+", " ", s).strip()
            if _SQL_START.match(flat):
                statements.append({"sql": flat, "op": flat.split(" ", 1)[0].lower()})
        connections = sorted(set(re.findall(r"getConnection\(\s*GlobalVariable\.(\w+)", body)))
        methods.append({
            "name": m.group("name"),
            "keyword": bool(m.group("kw")),
            "args": [a.strip() for a in m.group("args").split(",") if a.strip()],
            "line": text.count("\n", 0, m.start()) + 1,
            "statements": statements,
            "connections": connections,
            # Other methods of this class it calls: their databases count as touched too.
            "calls": sorted(n for n in names if n != m.group("name") and re.search(rf"\b{n}\s*\(", body)),
            "swallows_sql_errors": bool(re.search(r"catch\s*\(\s*SQLException\s+\w+\s*\)\s*\{\s*\w+\.printStackTrace\(\)", body)),
            "interpolated_sql": any("${" in s["sql"] or "'\"+" in s["sql"] for s in statements),
            "count_equality_without_zero_guard": bool(
                re.search(r"if\s*\(\s*count\w*\s*==\s*count\w*\s*\)", body)
                and not re.search(r"count\w*\s*(?:!=|>)\s*['\"]?0", body)
            ),
            "uses_execute_query_for_dml": bool(
                any(s["op"] != "select" for s in statements) and "executeQuery(" in body and "executeUpdate(" not in body
            ),
        })
    return {"file": rel_path, "class": qualified, "methods": methods}
