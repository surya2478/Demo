"""Outputs: the Automation KB files, a draft change set for them, and downloads.

The KB files follow the Automation KB's own layout (app_profile.json,
components.json, flows/*.json), so kb_compose and the Knowledge Base page read
them as they read a crawled application. Components carry the 4-field steps
kb_compose expects plus `step_meta` with what Katalon adds (conditions,
frames, database steps, value sources).

Nothing is written to the KB here: `draft_change_set` records a draft that a
person approves on the Knowledge Base page (or from the import page, with
the approve permission).
"""
from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime
from typing import Any, Optional

from . import analysis, generate
from .model import slug

KB_TYPE = {"assert_visible": "assert", "assert_present": "assert"}


def _signature(comp: dict) -> str:
    return json.dumps([(s["type"], s.get("target"), (s.get("value") or {}).get("template")) for s in comp["steps"]])


def _sensitive_columns(tables: list[dict]) -> set[str]:
    out = set()
    for t in tables:
        for col in t["columns"]:
            values = [r.get(col, "") for r in t["rows"] if r.get(col)]
            if values and sum(1 for v in values if re.fullmatch(r"\d{9,22}", v)) >= max(1, len(values) * 0.8):
                out.add(col)
    return out


def kb_files(models: list[dict], mds: list[dict], profile: Optional[dict], tables: list[dict], *,
             app_key: str, app_name: str) -> dict[str, Any]:
    now = datetime.utcnow().strftime("%Y-%m-%d")
    sensitive = _sensitive_columns(tables)
    components: dict[str, dict] = {}
    signatures: dict[str, str] = {}
    flows: dict[str, dict] = {}
    columns_used: set[str] = set()
    for mod, md in zip(models, mds):
        renamed: dict[str, str] = {}
        tc_slug = slug((md.get("katalon_test_case") or md.get("title") or "tc").rsplit("/", 1)[-1])[:40]
        for comp in mod["components"]:
            cid = comp["id"]
            sig = _signature(comp)
            if cid in components and signatures[cid] != sig:
                cid = f"{cid}__{tc_slug}"
            renamed[comp["id"]] = cid
            if cid in components:
                components[cid]["_provenance"]["test_cases"].append(md.get("katalon_test_case"))
                continue
            signatures[cid] = sig
            conds = {c for s in comp["steps"] for c in s.get("conds") or []} | set(comp.get("when") or [])
            for g in comp.get("gaps") or []:
                conds |= set(g.get("conds") or [])
            params = sorted({p for s in comp["steps"] for p in re.findall(r"\{(\w+)\}", json.dumps(
                [s.get("action"), s.get("target"), (s.get("value") or {}).get("template")]))})
            columns_used.update(params)
            components[cid] = {
                "id": cid,
                "description": f"{comp['name']} -- from Katalon {(comp.get('katalon') or {}).get('type', 'step')} "
                               f"{(comp.get('katalon') or {}).get('ref', '')}".strip(),
                "params": params,
                "requires": comp["requires"],
                "leaves": comp["leaves"],
                "steps": [[s["action"], KB_TYPE.get(s["type"], s["type"]), s.get("target") or "", s.get("expected") or ""]
                          for s in comp["steps"]],
                "step_meta": [{k: s.get(k) for k in ("id", "source", "via", "conds", "object", "params", "note", "seconds",
                                                     "key", "keyword", "katalon_kind", "error_target", "capture")
                               if s.get(k) not in (None, [], {})}
                              | ({"value": (s.get("value") or {}).get("template")} if s.get("value") else {})
                              | ({"db": {"connections": s["db"]["connections"],
                                         "statements": [x["sql"] for x in s["db"]["statements"]],
                                         "require_rows": s["db"]["count_guard_missing"]}} if s.get("db") else {})
                              for s in comp["steps"]],
                "conditions": {c: mod["conditions"][c] for c in sorted(conds, key=lambda x: int(x[1:]))},
                "when": comp.get("when") or [],
                "gaps": comp.get("gaps") or [],
                "katalon": comp.get("katalon"),
                "_provenance": {"source": "katalon", "imported": now, "status": "unverified",
                                "test_cases": [md.get("katalon_test_case")],
                                "note": "Locators come from the QA team's Object Repository; not yet run against the live application."},
            }
        flow_id = f"flow_{tc_slug}"
        sequence = [{"component": renamed[f["component"]],
                     "when": [mod["conditions"][c]["text"] for c in f.get("when") or []]} for f in mod["flow"]]
        ui = [s for s in sequence if components[s["component"]]["requires"] != "any"]
        flows[f"flows/{flow_id}.json"] = {
            "flow_id": flow_id,
            "flow_name": md.get("title") or flow_id,
            "category": "Katalon",
            "subcategory": (md.get("katalon_test_case") or "").rsplit("/", 2)[-2] if md.get("katalon_test_case") else "",
            "_provenance": {"source": "katalon", "imported": now, "katalon_test_case": md.get("katalon_test_case"),
                            "script": md.get("script"), "export_format": md.get("format")},
            "target_match_keywords": sorted({w.lower() for w in re.findall(r"[A-Za-z]{4,}", md.get("title") or "")}),
            "components": sequence,
            "graph_edges": [{"from": a["component"], "to": b["component"], "action": b["component"]} for a, b in zip(ui, ui[1:])],
            "checkpoints": [s["id"] for c in mod["components"] for s in c["steps"] if s["type"] == "db_check"],
            "data_binding": {"sheet": md.get("data_sheet"), "key_column": (mod.get("data") or {}).get("key_column"),
                             "example_row": md.get("data_key")},
            "automation_boundary": {
                "automated_scope": " -> ".join(c["name"] for c in mod["components"]),
                "external_boundary": "Database preparation and checks run through the team's keywords (Katalon) or are gated by KATALON_DB_MODE (Playwright).",
            },
        }
    params: dict[str, dict] = {}
    for col in sorted(columns_used):
        values = sorted({r.get(col) for t in tables for r in t["rows"] if r.get(col)})
        if col in sensitive or not values:
            params[col] = {"type": "text", "_note": "Comes from the test data row; values are not stored in the KB."}
        else:
            params[col] = {"values": values[:200], "_column": col}

    variables = (profile or {}).get("variables") or []
    value = {v["name"]: v.get("value") for v in variables if v.get("kind") in ("url", "setting")}
    base_url = value.get("appUrl") or ""
    report_like = analysis.locator_issues
    defects = []
    for mod in models:
        for comp in mod["components"]:
            for s in comp["steps"]:
                for issue in report_like(s.get("target")) if s.get("target") else []:
                    if issue["severity"] != "info":
                        defects.append({"id": f"LOC-{slug(s.get('object') or s['id'])}-{issue['code']}",
                                        "what": f"{s.get('object') or s['id']}: {issue['text']}", "component": comp["id"]})
    unique_defects = list({d["id"]: d for d in defects}.values())
    capture = next((s for mod in models for c in mod["components"] for s in c["steps"] if s["type"] == "capture"), None)
    loader = next((s for mod in models for c in mod["components"] for s in c["steps"] if s["type"] == "wait_hidden" and s.get("note")), None)
    frames = sorted({s.get("target") for mod in models for c in mod["components"] for s in c["steps"] if s["type"] == "frame" and s.get("target")})
    profile_doc = {
        "app_id": app_key,
        "app_name": app_name,
        "base_url": base_url,
        "domain": re.sub(r"^https?://", "", base_url).split("/")[0],
        "framework": "Katalon-tested web application (steps and locators from the QA team's Katalon project)",
        "_provenance": {"source": "katalon", "imported": now, "method": "Katalon Import: converter export + Object Repository + keywords + profile. Not crawled."},
        "environments": {value.get("env") or "default": base_url} if base_url else {},
        "katalon": {
            "test_cases": [md.get("katalon_test_case") for md in mds],
            "data_sheets": sorted({md.get("data_sheet") for md in mds if md.get("data_sheet")}),
            "systems": [{"variable": v["name"], **(v.get("connection") or {})} for v in variables if v.get("kind") == "connection"],
            "settings": {k: v for k, v in value.items() if k != "appUrl"},
            "frames": frames,
            "loader": {"xpath": loader.get("target"), "timeout_s": loader.get("seconds")} if loader else None,
            "order_result": {"success_xpath": capture.get("target"), "error_xpath": capture.get("error_target"),
                             "request_id": "text after '='"} if capture else None,
        },
        "known_defects": unique_defects,
        "selectors": {},
    }
    components_doc = {
        "_about": ("Reusable business components imported from the QA team's Katalon tests (Katalon Import). "
                   "Each component is one called Katalon test case; `steps` are [action, type, target, expected] with "
                   "{column} placeholders for test data, {vault:x} for secrets and {env:x} for profile settings; "
                   "`step_meta` carries conditions, frames and database steps. Status: unverified until run."),
        "app_key": app_key,
        "pages": sorted({c["requires"] for c in components.values()} | {c["leaves"] for c in components.values()}),
        "params": params,
        "components": list(components.values()),
    }
    return {"app_profile.json": profile_doc, "components.json": components_doc, **flows}


# --- Change set ---------------------------------------------------------------------------

def draft_change_set(db, *, app_key: str, files: dict[str, Any], actor: Optional[str], title: str, import_id: str) -> int:
    from ..automation import knowledge_base as kb
    from ..journeys import changesets

    folder = kb.get_kb_root() / app_key
    folder.mkdir(parents=True, exist_ok=True)
    changes: list[dict] = []
    for rel, doc in files.items():
        path = folder / rel
        if not path.exists():
            changes.append({"file": rel, "op": "create", "path": changesets.ROOT, "value": doc,
                            "note": "New file from Katalon Import."})
            continue
        current = json.loads(path.read_text(encoding="utf-8"))
        if rel == "components.json":
            for comp in doc["components"]:
                changes.append({"file": rel, "op": "set", "path": f"components[id={comp['id']}]", "value": comp,
                                "note": f"Component {comp['id']} from Katalon."})
            for name, spec in doc["params"].items():
                changes.append({"file": rel, "op": "set", "path": f"params.{name}", "value": spec, "note": "Allowed values from the test data."})
            changes.append({"file": rel, "op": "set", "path": "pages", "value": sorted(set(current.get("pages") or []) | set(doc["pages"]))})
        elif rel == "app_profile.json":
            for key in ("katalon", "known_defects", "_provenance") + (() if current.get("base_url") else ("base_url", "domain", "environments")):
                changes.append({"file": rel, "op": "set", "path": key, "value": doc[key], "note": "From Katalon Import."})
        else:
            changes.append({"file": rel, "op": "delete", "path": changesets.ROOT})
            changes.append({"file": rel, "op": "create", "path": changesets.ROOT, "value": doc, "note": "Flow regenerated from Katalon."})
    cs = changesets.draft(db, app_id=app_key, kb="automation", source="qa", actor=actor, title=title,
                          description="Components, flow and profile from the QA team's Katalon project (Katalon Import). "
                                      "Locators are the team's Object Repository; not yet run against the application.",
                          changes=changes, origin={"kind": "katalon_import", "import": import_id})
    return cs.id


# --- Downloads ------------------------------------------------------------------------------

def _zip(entries: dict[str, str | bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in entries.items():
            z.writestr(name, content if isinstance(content, bytes) else content.encode("utf-8"))
    return buf.getvalue()


def kb_zip(files: dict[str, Any]) -> bytes:
    return _zip({name: json.dumps(doc, ensure_ascii=False, indent=2) for name, doc in files.items()})


def _file_id(case: dict) -> str:
    return re.sub(r"[^\w.-]+", "_", case["case_id"])[:80]


def scripts_zip(chosen: list[dict], models: list[dict], mds_meta: list[dict], frameworks: list[str]) -> tuple[bytes, list[dict]]:
    entries: dict[str, str] = {}
    checks = []
    if "playwright" in frameworks:
        entries["playwright/support/katalon-shim.ts"] = generate.PLAYWRIGHT_SHIM
        entries["playwright/support/db.ts"] = generate.PLAYWRIGHT_DB
        entries["playwright/playwright.config.ts"] = generate.PLAYWRIGHT_CONFIG
        entries["playwright/package.json"] = json.dumps({"name": "katalon-import-tests", "private": True,
                                                         "scripts": {"test": "playwright test"},
                                                         "devDependencies": {"@playwright/test": "^1.48.0"}}, indent=2)
        entries["playwright/README.md"] = PLAYWRIGHT_README
    for case in chosen:
        mod, meta = models[case["model"]], mds_meta[case["model"]]
        if "playwright" in frameworks:
            spec = generate.playwright_spec(case, mod, meta)
            entries[f"playwright/tests/{_file_id(case)}.spec.ts"] = spec
            checks.append({"case_id": case["case_id"], **generate.conformance(case, spec, "playwright")})
        if "katalon" in frameworks:
            script = generate.katalon_script(case, mod, meta)
            entries[f"katalon/{_file_id(case)}.groovy"] = script
            checks.append({"case_id": case["case_id"], **generate.conformance(case, script, "katalon")})
    entries["CONFORMANCE.json"] = json.dumps(checks, indent=1)
    return _zip(entries), checks


def cases_xlsx(chosen: list[dict], models: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Case ID", "Title", "Status", "Steps", "Problems"])
    for case in chosen:
        ws.append([case["case_id"], case["title"], case["status"], len(case["steps"]),
                   "\n".join(p["text"] for p in case["problems"])])
    steps = wb.create_sheet("Test cases")
    header = ["Case ID", "Title", "Step", "Component", "Condition", "Action", "Test data", "Expected result", "KB step", "Katalon step"]
    steps.append(header)
    for case in chosen:
        for r in generate.manual_rows(case, models[case["model"]]):
            steps.append([case["case_id"], case["title"], r["n"], r["component"], "\n".join(r["when"]), r["action"],
                          r["data"], r["expected"], r["key"], r["source"]])
    for sheet in (ws, steps):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="B71920")
        for col, width in zip("ABCDEFGHIJ", (16, 48, 8, 26, 34, 60, 24, 44, 28, 12)):
            sheet.column_dimensions[col].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def report_markdown(meta: dict, result: dict) -> str:
    a = result["analysis"]
    lines = [f"# Katalon Import report: {meta['name']}", "",
             f"Generated {result['generated_at']}. Files: {len(meta.get('files') or [])}.", "", "## Inputs", ""]
    for item in a["checklist"]:
        lines.append(f"- [{'x' if item['present'] else ' '}] {item['label']} - {item['why']}")
    lines += ["", "## Summary", "", "```", json.dumps(a["summary"], indent=1), "```", "", "## Findings", ""]
    for f in a["findings"]:
        lines += [f"### {f['severity'].upper()} - {f['title']}", "", f["detail"], "", f"**Action:** {f['action']}", ""]
    lines += ["## Database steps", ""]
    for d in a["db_steps"]:
        systems = ", ".join(f"{s['variable']} ({s.get('host', '?')}/{s.get('service', '?')})" for s in d["systems"]) or "unknown"
        lines.append(f"- `{d['id']}` {d['type']} - {d['name']} - {systems}")
    lines += ["", "## Locators with issues", ""]
    for loc in a["locators"]:
        if loc["issues"]:
            lines.append(f"- `{loc['path']}` `{loc['xpath']}` - " + "; ".join(i["text"] for i in loc["issues"]))
    return "\n".join(lines) + "\n"


PLAYWRIGHT_README = """# Generated Playwright tests (Katalon Import)

Each spec is one data row of the source Katalon test case. Steps carry `[TC n]` markers that match the
manual test case step for step; regenerate from the Knowledge Base instead of editing a spec by hand.

    npm install
    npx playwright install chrome   # or set KATALON_BROWSER_CHANNEL
    KATALON_appUserName=... KATALON_appPassword=... npx playwright test

Environment:

- `KATALON_<name>`: secrets and settings from the Katalon profile (e.g. KATALON_appPassword). Never written into the specs.
- `KATALON_DB_MODE`: `off` (default) records database steps as annotations; `check` runs read-only checks;
  `all` also runs data preparation. Implement `run()` in support/db.ts with your database driver first.
- `KATALON_COND_TIMEOUT_MS`: how long "if visible" checks wait (default 2000).
"""
