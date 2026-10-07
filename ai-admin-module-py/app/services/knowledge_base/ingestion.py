from __future__ import annotations

import io
import json
import logging
from typing import Iterator
import requests
from docx import Document

from ...config import settings
from ..automation import extraction
from . import postgres_client
from . import squads

logger = logging.getLogger(__name__)

def validate_document(document_name: str, content: bytes) -> str:
    file_type = extraction.type_for_filename(document_name)
    if file_type not in {"docx", "pdf", "xlsx"}:
        raise ValueError("Only DOCX, PDF and XLSX files are supported.")
    if not content:
        raise ValueError("The uploaded file is empty.")
    if len(content) > extraction.MAX_EXTRACT_BYTES:
        raise ValueError("The uploaded file exceeds the 25 MB limit.")
    try:
        extraction.validate_signature(file_type, content[:16])
    except extraction.ExtractionError as exc:
        raise ValueError(str(exc)) from exc
    return file_type


def _extract_document_text(document_name: str, content: bytes) -> str:
    file_type = validate_document(document_name, content)
    try:
        if file_type == "docx":
            return _extract_docx_text(content)
        if file_type == "xlsx":
            rows = extraction.extract_rows(content, file_type)
            return "\n\n".join(json.dumps(row, ensure_ascii=False) for row in rows)
        text, _pages = extraction.extract_text(content, file_type)
        if not text.strip():
            raise ValueError("No readable PDF text was found. Scanned PDFs require OCR before ingestion.")
        return text
    except extraction.ExtractionError as exc:
        raise ValueError(str(exc)) from exc

def _extract_docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    blocks = []
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split()).strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()).strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                blocks.append(row_text)
    return "\n\n".join(blocks).strip()


def _chunk_text(text: str, chunk_size: int = 1500, overlap: int = 250) -> list[str]:
    words = text.split()
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = start
        length = 0
        while end < len(words) and length + len(words[end]) + (1 if end > start else 0) <= chunk_size:
            length += len(words[end]) + (1 if end > start else 0)
            end += 1
        if end == start:
            end += 1
        chunks.append(" ".join(words[start:end]))
        if end >= len(words):
            break
        overlap_words = max(1, min(end - start - 1, overlap // 6))
        start = end - overlap_words
    return chunks

def _normalize_systems(systems: list[str] | None) -> list[str]:
    normalized: list[str] = []
    seen: set[str] = set()
    for raw in systems or []:
        name = " ".join(str(raw).split())
        if not name:
            continue
        if len(name) > 100:
            raise ValueError("System names must be 100 characters or fewer.")
        key = name.casefold()
        if key not in seen:
            normalized.append(name)
            seen.add(key)
    if len(normalized) > 50:
        raise ValueError("A maximum of 50 systems can be associated with one document.")
    return normalized

def _openai_embeddings(texts: list[str]) -> list[list[float]]:
    if not settings.SQUAD_KB_EMBEDDING_BASE_URL.strip():
        raise RuntimeError("SQUAD_KB_EMBEDDING_BASE_URL is not configured")
    headers = {"Content-Type": "application/json"}
    if settings.SQUAD_KB_EMBEDDING_API_KEY.strip():
        headers["Authorization"] = f"Bearer {settings.SQUAD_KB_EMBEDDING_API_KEY.strip()}"
    response = requests.post(
        f"{settings.SQUAD_KB_EMBEDDING_BASE_URL.strip().rstrip('/')}/embeddings",
        headers=headers,
        json={
            "model": settings.SQUAD_KB_EMBEDDING_MODEL,
            "input": texts,
            # "dimensions": settings.SQUAD_KB_EMBEDDING_DIMENSIONS,
        },
        timeout=settings.SQUAD_KB_EMBEDDING_TIMEOUT_SECONDS,
    )
    if response.status_code >= 400:
        raise RuntimeError(f"Embedding request failed: {response.text[:500]}")
    payload = response.json()
    return [item["embedding"] for item in sorted(payload.get("data", []), key=lambda item: item["index"])]


def ingest_docx_progress(
    squad_name: str, document_name: str, content: bytes, systems: list[str] | None = None
) -> Iterator[dict]:
    """Ingest DOCX, PDF or XLSX with extraction, embedding and storage progress."""
    squad_name = squads.normalize_squad_name(squad_name)
    if not squad_name:
        raise ValueError("Unsupported Squad Knowledge Base")
    systems = _normalize_systems(systems)

    yield {"type": "progress", "stage": "extract", "status": "running", "message": f"Extracting text from {document_name}."}
    text = _extract_document_text(document_name, content)
    if not text:
        raise ValueError("The document contains no readable text or populated spreadsheet rows.")
    yield {
        "type": "progress",
        "stage": "extract",
        "status": "completed",
        "message": f"Extracted {len(text):,} characters.",
        "characters": len(text),
    }

    chunks = _chunk_text(text)
    yield {
        "type": "progress",
        "stage": "embedding",
        "status": "running",
        "message": f"Creating embeddings for {len(chunks)} text chunks.",
        "total_chunks": len(chunks),
    }
    embeddings = _openai_embeddings(chunks)
    if len(embeddings) != len(chunks):
        raise RuntimeError("OpenAI returned an incomplete embedding batch")
    yield {
        "type": "progress",
        "stage": "embedding",
        "status": "completed",
        "message": f"Created {len(embeddings)} embeddings.",
        "total_chunks": len(chunks),
    }

    target_squads = ("B2B", "BCRM", "SMB") if squad_name == "Device Squad" else (squad_name,)
    yield {
        "type": "progress",
        "stage": "store",
        "status": "running",
        "message": f"Storing document in {squad_name}.",
        "systems": systems,
        "completed_chunks": 0,
        "total_chunks": len(chunks),
    }
    for index, (chunk, embedding) in enumerate(zip(chunks, embeddings), start=1):
        for target_squad in target_squads:
            postgres_client.save_kb_entry(target_squad, document_name, chunk, embedding, systems)

        yield {
            "type": "progress",
            "stage": "store",
            "status": "running",
            "message": f"Stored {index} of {len(chunks)} chunks.",
            "completed_chunks": index,
            "total_chunks": len(chunks),
        }

    yield {
        "type": "progress",
        "stage": "store",
        "status": "completed",
        "message": f"Stored {len(chunks)} chunks in {squad_name}.",
        "completed_chunks": len(chunks),
        "total_chunks": len(chunks),
    }
    yield {
        "type": "result",
        "data": {
            "document_name": document_name,
            "squad_name": squad_name,
            "systems": systems,
            "chunks_ingested": len(chunks),
            "embedding_model": settings.SQUAD_KB_EMBEDDING_MODEL,
        },
    }


def ingest_docx(
    squad_name: str,
    document_name: str,
    content: bytes,
    systems: list[str] | None = None,
) -> dict[str, int | str | list[str]]:
    """Ingest a document and return its result, for non-streaming callers."""
    result = None
    for event in ingest_docx_progress(squad_name, document_name, content, systems):
        if event["type"] == "result":
            result = event["data"]
    if result is None:
        raise RuntimeError("Document ingestion completed without a result")
    return result

# def ingest_docx(squad_name: str, document_name: str, content: bytes) -> dict[str, int | str]:
#     if squad_name not in {"Device Squad", "B2B", "BCRM", "SMB"}:
#         raise ValueError("Unsupported Squad Knowledge Base")
#     text = _extract_docx_text(content)
#     if not text:
#         raise ValueError("The DOCX contains no readable text")

#     chunks = _chunk_text(text)
#     embeddings = _openai_embeddings(chunks)
#     if len(embeddings) != len(chunks):
#         raise RuntimeError("OpenAI returned an incomplete embedding batch")

#     target_squads = ("B2B", "BCRM", "SMB") if squad_name == "Device Squad" else (squad_name,)
#     for chunk, embedding in zip(chunks, embeddings):
#         for target_squad in target_squads:
#             postgres_client.save_kb_entry(target_squad, document_name, chunk, embedding)

#     return {
#         "document_name": document_name,
#         "squad_name": squad_name,
#         "chunks_ingested": len(chunks),
#         "embedding_model": settings.SQUAD_KB_EMBEDDING_MODEL,
#     }