"""
--- Automation Module: document intake smoke test ----------------------------

Drives the /automation/intake/* endpoints over real HTTP with real DOCX, XLSX
and CSV files built on the fly. Nothing outside localhost is contacted, and no
model is called: extraction reads a clean table directly and only falls back to
the model when it cannot, so every document here takes the deterministic path.
The model fallback is therefore NOT covered by this test.

What it is really checking is the claim the feature rests on: that a test case
document lands rows in automation_test_cases with the same shape and the same
eligibility scoring as a case that arrived through the review workflow, so
everything downstream cannot tell the two apart.

Auth is handled by overriding get_current_user_strict, the dependency
require_permission nests. The permission check still runs against the injected
user's roles.

    .venv312/Scripts/python.exe tests/automation/smoke_intake_api.py
"""
import io
import sys
import threading
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import requests  # noqa: E402
import uvicorn  # noqa: E402

from app.main import app  # noqa: E402
from app.database.connection import get_db  # noqa: E402
from app.models.automation import (  # noqa: E402
    AutomationIntakeBatch,
    AutomationTestCase,
)
from app.utils.security import get_current_user_strict  # noqa: E402

PORT = 8046
BASE = f"http://127.0.0.1:{PORT}/automation/intake"
BATCH_NAME = "SMOKE-INTAKE"

FAKE_ADMIN = {
    "username": "smoke-intake",
    "role": "admin",
    "roles": ["admin"],
    "permissions": ["automation.view", "automation.execute", "automation.approve"],
}

app.dependency_overrides[get_current_user_strict] = lambda: FAKE_ADMIN

failures: list[str] = []


def check(label: str, condition: bool, detail: str = "") -> None:
    print(f"  {'PASS' if condition else 'FAIL'}  {label}")
    if not condition:
        failures.append(f"{label} {detail}".strip())


def serve() -> uvicorn.Server:
    config = uvicorn.Config(app, host="127.0.0.1", port=PORT, log_level="error")
    server = uvicorn.Server(config)
    threading.Thread(target=server.run, daemon=True).start()
    for _ in range(100):
        if server.started:
            return server
        time.sleep(0.1)
    raise RuntimeError("the test server did not start")


def build_xlsx() -> bytes:
    import pandas as pd

    buffer = io.BytesIO()
    pd.DataFrame(
        [
            {
                "Test Case ID": "SMK-101",
                "Category": "Authentication",
                "Test Case Description": "Verify a registered user can sign in",
                "Preconditions": "An active account exists",
                "Test Steps": "1. Open the login page. 2. Enter valid credentials. 3. Click Sign In.",
                "Test Data": "username=smoke.user",
                "Expected Result": "The dashboard is displayed",
            },
            {
                "Test Case ID": "SMK-102",
                "Category": "Authentication",
                "Test Case Description": "Verify an invalid password is rejected",
                "Preconditions": "An active account exists",
                "Test Steps": "1. Open the login page. 2. Enter a wrong password. 3. Click Sign In.",
                "Test Data": "password=wrong",
                "Expected Result": "An invalid credentials message is shown",
            },
        ]
    ).to_excel(buffer, index=False)
    return buffer.getvalue()


def build_brd_docx() -> bytes:
    """A BRD whose substance is in a table -- the case naive extractors miss."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Business Requirements Document: Sign-in")
    doc.add_paragraph("The system shall lock an account after three failed attempts.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Requirement"
    table.rows[0].cells[1].text = "Priority"
    table.rows[1].cells[0].text = "Account lockout after three failures"
    table.rows[1].cells[1].text = "High"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def cleanup(db) -> None:
    for batch in db.query(AutomationIntakeBatch).filter(
        AutomationIntakeBatch.name == BATCH_NAME
    ).all():
        db.delete(batch)
    db.query(AutomationTestCase).filter(
        AutomationTestCase.display_id.in_(["SMK-101", "SMK-102", "SMK-201"])
    ).delete(synchronize_session=False)
    db.commit()


def main() -> int:
    server = serve()
    db = next(get_db())
    batch_id = None

    try:
        cleanup(db)

        print("\nBatch lifecycle")
        resp = requests.post(f"{BASE}/batches", json={"name": BATCH_NAME}, timeout=30)
        check("creates a batch", resp.status_code == 201, f"got {resp.status_code}: {resp.text[:160]}")
        if resp.status_code != 201:
            return 1
        batch_id = resp.json()["id"]
        check("starts empty", resp.json()["document_count"] == 0)

        resp = requests.post(f"{BASE}/batches", json={"name": ""}, timeout=30)
        check("refuses a blank name", resp.status_code == 422, f"got {resp.status_code}")

        print("\nTest case document -> candidates")
        resp = requests.post(
            f"{BASE}/batches/{batch_id}/documents",
            data={"role": "test_cases"},
            files={"files": ("Test_Cases.xlsx", build_xlsx(),
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=60,
        )
        check("accepts the upload", resp.status_code == 201, f"got {resp.status_code}: {resp.text[:200]}")
        batch = resp.json()["batch"]
        doc = batch["documents"][0]
        check("extracts the text", doc["extract_status"] == "ready", doc.get("extract_error") or "")
        check(
            "uploading does NOT create test cases on its own",
            doc["cases_found"] == 0,
            f"found {doc['cases_found']}",
        )

        print("\nExtract Test Cases")
        # A clean spreadsheet is read from its own table, so this path never
        # calls the model and runs with no provider configured. The model
        # fallback for documents with no readable table is not covered here.
        resp = requests.post(f"{BASE}/batches/{batch_id}/extract", timeout=30)
        check("returns 202 with a job", resp.status_code == 202, f"got {resp.status_code}: {resp.text[:200]}")
        job_id = resp.json().get("job_id")
        check("carries a job id", bool(job_id))

        settled = None
        for _ in range(60):
            j = requests.get(
                f"http://127.0.0.1:{PORT}/automation/workbench/jobs/{job_id}", timeout=30
            ).json()
            if j["status"] in ("completed", "failed"):
                settled = j
                break
            time.sleep(0.4)
        check("the job settles", settled is not None, "timed out")
        if settled:
            check("and completes", settled["status"] == "completed", settled.get("message", ""))
            check("reporting both cases", "2 test case" in settled.get("message", ""), settled.get("message", ""))

        resp = requests.get(f"{BASE}/batches", timeout=30)
        after = next(b for b in resp.json()["batches"] if b["id"] == batch_id)
        doc = after["documents"][0]
        check("the document now shows its count", doc["cases_found"] == 2, f"got {doc['cases_found']}")
        check("and the batch total rolls up", after["cases_found"] == 2)

        print("\nRe-extracting is idempotent")
        resp = requests.post(f"{BASE}/batches/{batch_id}/extract", timeout=30)
        job_id = resp.json()["job_id"]
        for _ in range(60):
            j = requests.get(
                f"http://127.0.0.1:{PORT}/automation/workbench/jobs/{job_id}", timeout=30
            ).json()
            if j["status"] in ("completed", "failed"):
                break
            time.sleep(0.4)
        db.expire_all()
        check(
            "a second extraction adds no duplicates",
            db.query(AutomationTestCase)
            .filter(AutomationTestCase.display_id.in_(["SMK-101", "SMK-102"]))
            .count()
            == 2,
        )

        rows = (
            db.query(AutomationTestCase)
            .filter(AutomationTestCase.display_id.in_(["SMK-101", "SMK-102"]))
            .order_by(AutomationTestCase.display_id)
            .all()
        )
        check("wrote both rows to automation_test_cases", len(rows) == 2, f"got {len(rows)}")
        if rows:
            first = rows[0]
            check("mapped the steps column", "Enter valid credentials" in (first.steps_raw or ""))
            check("mapped the expected result", "dashboard" in (first.expected_result or "").lower())
            check(
                "scored eligibility like any other case",
                first.eligibility_verdict in ("ready", "review", "needs_enhancement")
                and first.eligibility_score > 0,
                f"{first.eligibility_verdict} / {first.eligibility_score}",
            )

        print("\nHarvested cases reach the Workbench (the second way in)")
        # The whole point of the feature: a case read out of a document must
        # show up in "What to refine" beside one that arrived by approval, and
        # be tagged so the UI knows which request field its id belongs in.
        resp = requests.get(
            f"http://127.0.0.1:{PORT}/automation/workbench/sources", timeout=30
        )
        check("workbench sources load", resp.status_code == 200, f"got {resp.status_code}")
        listed = resp.json()["sources"]
        harvested = [s for s in listed if s["tc_display_id"] in ("SMK-101", "SMK-102")]
        check("both harvested cases are listed", len(harvested) == 2, f"got {len(harvested)}")
        if harvested:
            one = harvested[0]
            check("tagged origin=studio", one["origin"] == "studio", str(one.get("origin")))
            check("carries a title", bool(one["title"]))
            check(
                "carries full eligibility, not a partial dict",
                {"score", "verdict", "reasons", "step_count", "can_generate"}
                <= set(one["eligibility"]),
                str(sorted(one["eligibility"])),
            )
            check("not yet refined", one["refined_test_case_id"] is None)
            check(
                "its id is the studio primary key, usable as case_ids",
                str(one["id"]).isdigit(),
                str(one["id"]),
            )
        check(
            "approved-origin rows still carry origin too",
            all("origin" in s for s in listed),
        )

        print("\nBRD document -> text only, no candidates")
        resp = requests.post(
            f"{BASE}/batches/{batch_id}/documents",
            data={"role": "brd"},
            files={"files": ("Test_BRD.docx", build_brd_docx(),
                             "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            timeout=60,
        )
        check("accepts the BRD", resp.status_code == 201, f"got {resp.status_code}")
        brd = next(d for d in resp.json()["batch"]["documents"] if d["filename"] == "Test_BRD.docx")
        check("extracts the BRD text", brd["extract_status"] == "ready", brd.get("extract_error") or "")
        check("reads the DOCX table, not just paragraphs", brd["text_chars"] > 120, f"{brd['text_chars']} chars")
        check("produces no test cases from a BRD", brd["cases_found"] == 0)

        print("\nRejections and failures are legible")
        resp = requests.post(
            f"{BASE}/batches/{batch_id}/documents",
            data={"role": "brd"},
            files={"files": ("notes.rtf", b"whatever", "application/rtf")},
            timeout=30,
        )
        rejected = resp.json()["rejected"]
        check("rejects an unsupported extension", len(rejected) == 1, str(rejected))
        check("and names the supported formats", "xlsx" in (rejected[0]["reason"] if rejected else ""))

        resp = requests.post(
            f"{BASE}/batches/{batch_id}/documents",
            data={"role": "brd"},
            files={"files": ("fake.pdf", b"this is not a pdf at all", "application/pdf")},
            timeout=30,
        )
        bad = next(d for d in resp.json()["batch"]["documents"] if d["filename"] == "fake.pdf")
        check("records a signature mismatch as failed", bad["extract_status"] == "failed")
        check("and says why", "not a PDF" in (bad["extract_error"] or ""), bad.get("extract_error") or "")

        print("\nRole correction relabels; extraction stays explicit")
        resp = requests.post(
            f"{BASE}/batches/{batch_id}/documents",
            data={"role": "brd"},  # deliberately wrong
            files={"files": ("More_Cases.csv",
                             b"Test Case ID,Test Steps,Expected Result\n"
                             b"SMK-201,1. Open it. 2. Click go.,It works\n", "text/csv")},
            timeout=30,
        )
        mis = next(d for d in resp.json()["batch"]["documents"] if d["filename"] == "More_Cases.csv")
        check("no cases while mislabelled as BRD", mis["cases_found"] == 0)

        resp = requests.patch(f"{BASE}/documents/{mis['id']}", json={"role": "test_cases"}, timeout=30)
        check("role change succeeds", resp.status_code == 200, f"got {resp.status_code}: {resp.text[:160]}")
        check(
            "relabelling alone spends no model call and creates nothing",
            next(d for d in resp.json()["batch"]["documents"] if d["filename"] == "More_Cases.csv")["cases_found"] == 0,
        )
        check("and it says what to press next", "Extract" in (resp.json().get("note") or ""), str(resp.json().get("note")))

        resp = requests.post(f"{BASE}/batches/{batch_id}/extract", timeout=30)
        job_id = resp.json()["job_id"]
        for _ in range(60):
            j = requests.get(
                f"http://127.0.0.1:{PORT}/automation/workbench/jobs/{job_id}", timeout=30
            ).json()
            if j["status"] in ("completed", "failed"):
                break
            time.sleep(0.4)
        db.expire_all()
        check(
            "extracting after the correction reads it in",
            db.query(AutomationTestCase).filter(AutomationTestCase.display_id == "SMK-201").count() == 1,
        )

        print("\nDeletion")
        resp = requests.delete(f"{BASE}/documents/{brd['id']}", timeout=30)
        check("detaches a document", resp.status_code == 200, f"got {resp.status_code}")
        check(
            "and removes it from the batch",
            all(d["filename"] != "Test_BRD.docx" for d in resp.json()["batch"]["documents"]),
        )

        resp = requests.delete(f"{BASE}/batches/{batch_id}", timeout=30)
        check("deletes the batch", resp.status_code == 200, f"got {resp.status_code}")
        db.expire_all()
        check(
            "harvested test cases survive the batch deletion",
            db.query(AutomationTestCase).filter(AutomationTestCase.display_id == "SMK-101").count() == 1,
        )

        print("\nListing")
        resp = requests.get(f"{BASE}/batches", timeout=30)
        check("lists batches", resp.status_code == 200, f"got {resp.status_code}")
        check(
            "and the deleted one is gone",
            all(b["name"] != BATCH_NAME for b in resp.json()["batches"]),
        )
    finally:
        cleanup(db)
        db.close()
        server.should_exit = True

    print()
    if failures:
        print(f"{len(failures)} check(s) failed:")
        for failure in failures:
            print(f"  - {failure}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
