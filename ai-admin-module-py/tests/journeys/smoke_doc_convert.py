"""--- Screenshots on Azure, 30 MB documents, PDF to structured Word ----------------------

  1. With extraction on Azure, the Azure deployment reads the screenshots;
     JOURNEY_EXTRACTION_VISION=off reads text only.
  2. A tall Figma page is sent as several slices, a normal page as one.
  3. Extraction sends page text and slices to Azure, two pages per call.
  4. The upload limit is 30 MB.
  5. Convert to Word, AI-structured: a heading per page, the screens table, the
     UI elements table and page pictures; then text only; download; add the
     .docx as a document of its own; deleting the PDF removes its conversion.

Run from ai-admin-module-py:

    .venv312/Scripts/python.exe tests/journeys/smoke_doc_convert.py
"""
import io
import os
import shutil
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
PDF = Path(r"C:\Test_AI_Agents\AI-Testing-Review\KB-Claude\Sign-in using Password.pdf")
if not PDF.exists():
    print(f"SKIPPED: {PDF} is not on this machine.")
    sys.exit(0)

work = Path(tempfile.mkdtemp())
shutil.copytree(ROOT / "knowledge_base" / "journeys", work / "journeys")
os.environ["JOURNEY_KB_PATH"] = str(work / "journeys")
os.environ["DATABASE_URL"] = f"sqlite:///{(work / 'kb.db').as_posix()}"

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.config import settings  # noqa: E402
from app.database.connection import Base, SessionLocal, engine  # noqa: E402
import app.models.automation  # noqa: E402,F401
import app.models.kb_changes  # noqa: E402,F401
from app.models.kb_changes import KbSourceDocument  # noqa: E402
from app.routers import knowledge_base_page  # noqa: E402
from app.services.automation import llm  # noqa: E402
from app.services.generation import azure_llm  # noqa: E402
from app.services.journeys import doc_convert, documents, store  # noqa: E402
from app.utils.security import get_current_user_strict  # noqa: E402
import app.services.authz as authz  # noqa: E402

Base.metadata.create_all(bind=engine)
APP = "eand_b2b_portal"
failures = []


def check(condition, message):
    print(("  ok   " if condition else "  FAIL ") + message)
    if not condition:
        failures.append(message)


# Azure "configured": the deployment name is all the code asks for here.
azure_llm.missing_settings = lambda: []
azure_llm.model_name = lambda: "gpt-5.1"
settings.JOURNEY_EXTRACTION_PROVIDER = "azure"
settings.JOURNEY_EXTRACTION_VISION = "on"

calls = []


def fake_images(*, system_prompt, user_input, images, model, model_cls, provider=None, **kwargs):
    calls.append({"provider": provider, "images": len(images), "user": user_input, "cls": model_cls.__name__,
                  "jpeg": all(i[:3] == b"\xff\xd8\xff" for i in images)})
    if model_cls is doc_convert.CPage:
        return doc_convert.CPage(
            screen_name="Sign In", purpose="The user signs in with a username and password.",
            sections=[doc_convert.CSection(heading="Login form", lines=["Username field", "Password field", "SIGN IN button"])],
            elements=[doc_convert.CElement(label="Username or email*", type="text field", notes="mandatory"),
                      doc_convert.CElement(label="SIGN IN", type="button")],
            actions=[doc_convert.CAction(trigger="Click SIGN IN", result="Select a Organization to continue")],
            notes=["Password is masked"])
    return model_cls()


llm.complete_json_with_images = fake_images
llm.complete_json = lambda **kwargs: (_ for _ in ()).throw(AssertionError("text-only call while vision is on"))

api = FastAPI()
api.include_router(knowledge_base_page.router, prefix="/kb")
api.dependency_overrides[get_current_user_strict] = lambda: {"username": "analyst", "roles": ["tester"], "permissions": []}
authz.has_permission = lambda roles, perm: perm in ("automation.view", "automation.execute")
client = TestClient(api)
kb = store.load(APP)

print("1. Who reads the screenshots")
check(documents.vision_model() == "azure:gpt-5.1", f"Azure deployment ({documents.vision_model()})")
settings.JOURNEY_EXTRACTION_VISION = "off"
check(documents.vision_model() is None, "JOURNEY_EXTRACTION_VISION=off: text only")
settings.JOURNEY_EXTRACTION_VISION = "on"
check(client.get(f"/kb/{APP}/documents").json()["vision_model"] == "azure:gpt-5.1", "the page is told who reads them")
body = azure_llm._chat_body("sys", "hello", 100, None, [b"\xff\xd8\xff\xe0jpeg", b"\x89PNG"])
parts = body["messages"][1]["content"]
check(parts[0] == {"type": "text", "text": "hello"} and parts[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")
      and parts[2]["image_url"]["url"].startswith("data:image/png;base64,"), "Azure request carries the text then the pictures")
check(azure_llm._chat_body("sys", "hello", 100, None)["messages"][1]["content"] == "hello", "no pictures: plain text as before")

print("2. Tall pages in slices")
import pymupdf  # noqa: E402
tall = pymupdf.open()
page = tall.new_page(width=1440, height=9000)
for y in range(200, 9000, 600):
    page.insert_text((100, y), f"Figma frame label at {y}", fontsize=40)
tall.new_page(width=612, height=792).insert_text((72, 72), "A normal page")
tall_bytes = tall.tobytes()
r = client.post(f"/kb/{APP}/documents", files={"file": ("Figma flow.pdf", tall_bytes, "application/pdf")}, data={"doc_kind": "screen_flow"})
check(r.status_code == 201, f"tall PDF uploaded ({r.status_code})")
tall_doc = SessionLocal().get(KbSourceDocument, r.json()["id"])
slices = documents.page_images(kb, tall_doc, 1)
check(len(slices) == 4 and all(s[:3] == b"\xff\xd8\xff" for s in slices), f"1440x9000 page -> 4 JPEG slices ({len(slices)})")
check(len(documents.page_images(kb, tall_doc, 2)) == 1, "a normal page -> one picture")

print("3. Extraction on Azure with screenshots")
r = client.post(f"/kb/{APP}/documents", files={"file": (PDF.name, PDF.read_bytes(), "application/pdf")}, data={"doc_kind": "screen_flow"})
pdf_id = r.json()["id"]
calls.clear()
client.post(f"/kb/documents/{pdf_id}/extract")
for _ in range(100):
    doc = client.get(f"/kb/documents/{pdf_id}").json()
    if doc["status"] != "extracting":
        break
    time.sleep(0.1)
check(doc["status"] == "extracted", f"extracted ({doc['status']}: {doc.get('error')})")
check(calls and all(c["provider"] == "azure" and c["jpeg"] for c in calls), "every call went to Azure with JPEG pictures")
check(len(calls) == 1 and calls[0]["images"] == 2, f"2 pages in one call, one picture each ({[c['images'] for c in calls]})")
check("several slices" in calls[0]["user"], "the model is told a tall page comes in slices")
check(doc["extraction"]["vision"] and doc["model"] == "azure:gpt-5.1", f"extraction records vision by {doc['model']}")

print("3b. An extraction cut off by a backend restart")
db = SessionLocal()
row = db.get(KbSourceDocument, pdf_id)
row.status = "extracting"  # as left behind when the thread died
db.commit()
db.close()
listed = next(d for d in client.get(f"/kb/{APP}/documents").json()["documents"] if d["id"] == pdf_id)
check(listed["status"] == "failed" and "restarted" in (listed["error"] or ""), f"marked failed, not stuck ({listed['status']}: {listed['error']})")
calls.clear()
client.post(f"/kb/documents/{pdf_id}/extract")
for _ in range(100):
    doc = client.get(f"/kb/documents/{pdf_id}").json()
    if doc["status"] != "extracting":
        break
    time.sleep(0.1)
check(doc["status"] == "extracted", f"Extract again works ({doc['status']})")
documents._active.add(pdf_id)
db = SessionLocal()
db.get(KbSourceDocument, pdf_id).status = "extracting"
db.commit()
db.close()
check(client.get(f"/kb/documents/{pdf_id}").json()["status"] == "extracting", "a live extraction is left alone")
documents._active.discard(pdf_id)
db = SessionLocal()
db.get(KbSourceDocument, pdf_id).status = "extracted"
db.commit()
db.close()

print("4. 30 MB")
check(documents.MAX_BYTES == 30 * 1024 * 1024, "limit is 30 MB")
r = client.post(f"/kb/{APP}/documents", files={"file": ("big.pdf", b"%PDF-1.7\n" + b"0" * (30 * 1024 * 1024), "application/pdf")})
check(r.status_code == 400 and "30 MB" in r.json()["detail"], f"over 30 MB refused ({r.status_code} {r.json().get('detail')})")

print("5. Convert to Word")
from docx import Document  # noqa: E402


def wait(doc_id):
    for _ in range(200):
        state = client.get(f"/kb/documents/{doc_id}/convert").json()
        if state["status"] != "running":
            return state
        time.sleep(0.1)
    return state


check(client.get(f"/kb/documents/{pdf_id}/convert").json()["status"] == "none", "nothing converted yet")
calls.clear()
r = client.post(f"/kb/documents/{pdf_id}/convert", json={"mode": "ai", "images": True})
check(r.status_code == 202 and r.json()["status"] == "running", f"started ({r.status_code})")
state = wait(pdf_id)
check(state["status"] == "done" and state["ready"] and state["page"] == 2 and state["pages"] == 2, f"done, page 2 of 2 ({state})")
check(len(calls) == 2 and all(c["cls"] == "CPage" and c["provider"] == "azure" for c in calls), "one Azure call per page")
listed = next(d for d in client.get(f"/kb/{APP}/documents").json()["documents"] if d["id"] == pdf_id)
check(listed["conversion"]["ready"], "the documents list shows the Word version")
r = client.get(f"/kb/documents/{pdf_id}/convert.docx")
check(r.status_code == 200 and "Sign-in%20using%20Password%20%28structured%29.docx" in r.headers["content-disposition"], "downloads as '<name> (structured).docx'")
word = Document(io.BytesIO(r.content))
headings = [p.text for p in word.paragraphs if p.style.name.startswith("Heading")]
check("Page 1 · Sign In" in headings and "Page 2 · Sign In" in headings, f"a heading per page ({headings[:4]})")
check({"UI elements", "Actions and navigation", "Notes", "Login form"} <= set(headings), "elements, actions, notes and sections")
cells = [c.text for t in word.tables for row in t.rows for c in row.cells]
check("Username or email*" in cells and "Screen" in cells, "screens table and UI elements table")
check(len(word.inline_shapes) == 2, f"a picture per page ({len(word.inline_shapes)})")
check(any("SIGN IN → Select a Organization to continue" in p.text for p in word.paragraphs), "actions read 'trigger -> result'")

r = client.post(f"/kb/documents/{tall_doc.id}/convert", json={"mode": "text", "images": True})
state = wait(tall_doc.id)
word = Document(io.BytesIO(client.get(f"/kb/documents/{tall_doc.id}/convert.docx").content))
check(state["status"] == "done" and any("Figma frame label at 200" in p.text for p in word.paragraphs), "text only: the page's own text")
check(len(word.inline_shapes) == 5, f"tall page as 4 picture slices + 1 ({len(word.inline_shapes)})")

r = client.post(f"/kb/documents/{pdf_id}/convert/add")
added = r.json()
check(r.status_code == 201 and added["filename"] == "Sign-in using Password (structured).docx" and added["doc_kind"] == "screen_flow",
      f"added as a document of its own ({r.status_code} {added.get('filename')})")
check(client.post(f"/kb/documents/{added['id']}/convert", json={"mode": "ai"}).status_code == 400, "a .docx cannot be converted again")
check(client.post(f"/kb/documents/{pdf_id}/convert/add").status_code == 200, "adding it twice keeps one copy")

files = list((work / "journeys" / APP / "sources" / "conversions").glob("*"))
check(len(files) == 4, f".docx and progress file for both PDFs ({len(files)})")
client.delete(f"/kb/documents/{tall_doc.id}")
left = list((work / "journeys" / APP / "sources" / "conversions").glob("*"))
check(len(left) == 2 and all("Figma" not in f.name for f in left), "deleting the PDF removes its conversion")

shutil.rmtree(work, ignore_errors=True)
if failures:
    print(f"\n{len(failures)} check(s) failed")
    sys.exit(1)
print("\nall checks passed")
